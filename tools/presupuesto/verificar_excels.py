"""Reproduce las comprobaciones del análisis docs/presupuesto/ANALISIS_EXCEL_PRESUPUESTO.md.

Uso:
    python3 tools/presupuesto/verificar_excels.py <ElSuperDocumento.xlsx> <PRESUPUESTO_2026.xlsx>

Requiere: pip install openpyxl
"""
import re
import sys
import warnings
from datetime import datetime

import openpyxl
from openpyxl.utils import column_index_from_string as ci

warnings.filterwarnings("ignore")


def num(v):
    return v if isinstance(v, (int, float)) else 0


def main(super_path, ppto_path):
    sup_f = openpyxl.load_workbook(super_path)
    sup_v = openpyxl.load_workbook(super_path, data_only=True)
    ppt_f = openpyxl.load_workbook(ppto_path)
    ppt_v = openpyxl.load_workbook(ppto_path, data_only=True)
    pv, pf = ppt_v["PPTO CON PROP"], ppt_f["PPTO CON PROP"]
    cfv, cff = ppt_v["CASH FLOW"], ppt_f["CASH FLOW"]
    prot_f = ppt_f["PROTEC CON PROP"]

    print("== 1. Totales del presupuesto (PPTO CON PROP) ==")
    for c, desc in [("G124", "Comprobación (debe ser 0)"), ("G122", "Total por naturaleza"),
                    ("G123", "Total por cuentas"), ("G103", "Personal coordinación (residual)"),
                    ("AO113", "Subvenciones"), ("AP113", "Convenios"),
                    ("AQ113", "Donaciones"), ("AR113", "Facturación")]:
        print(f"  {c:6} {desc:35} {pv[c].value:>16,.2f}")
    print("  Fórmula G103:", pf["G103"].value)
    prot = ppt_v["PROTEC CON PROP"]
    print(f"  PROTEC ingresos N9={prot['N9'].value:,.2f}  gastos K26={prot['K26'].value:,.2f}  "
          f"inversiones K31={prot['K31'].value:,.2f}")

    print("\n== 2. CASH FLOW ==")
    for c in ("AO116", "AP116", "AQ116", "AR116"):
        print(f"  {cfv[c.replace('116', '115')].value:20} {cfv[c].value:>16,.2f}")
    print(f"  Fila 89 (4101) AO={cfv['AO89'].value}  Fila 90 (4104) AO={cfv['AO90'].value}"
          "  <- importe idéntico, sospecha de copia")
    rangos = {}
    for col in range(7, 42):
        f = cff.cell(85, col).value
        m = re.search(r"SUM\(\w+?4:\w+?(\d+)\)", str(f))
        if m:
            rangos.setdefault(m.group(1), []).append(cff.cell(85, col).column_letter)
    print("  Fila 85 CASH FLOW: última fila sumada por columna ->",
          {k: f"{v[0]}..{v[-1]} ({len(v)} cols)" for k, v in rangos.items()})

    print("\n== 3. Reparto gastos comunes ==")
    print("  PPTO  G107:", pf["G107"].value, "| G109:", pf["G109"].value)
    print("  PROTEC D10:", prot_f["D10"].value, "| G10:", prot_f["G10"].value)

    print("\n== 4. Calidad de datos en PPTO CON PROP ==")
    codigos = {}
    for r in range(4, 100):
        cod = pv.cell(r, 3).value
        if cod:
            codigos.setdefault(str(cod), []).append(r)
        for col in (5, 6):
            v = pv.cell(r, col).value
            if v is not None and not isinstance(v, datetime):
                print(f"  Fecha no válida {pv.cell(r, col).coordinate}: {v!r} ({pv.cell(r, 2).value})")
    for cod, rows in codigos.items():
        if len(rows) > 1:
            print(f"  Código repetido {cod} en filas {rows}")
    print("  Cabecera GASTOS No Imputables O1:", ppt_v["GASTOS No Imputables"]["O1"].value)

    print("\n== 5. SuperDocumento: alineación de filas con PROYECTOS ==")
    for sh in ["PREV_PERSONAL", "PREV_MATERIAL", "EJE_PERSONAL", "EJE_MATERIAL",
               "control_plurianuales", "plurianuales_ingresos", "anuales_ingresos",
               "facturables", "INGRESOS TOTALES"]:
        ws = sup_f[sh]
        bad = 0
        for r in range(4, ws.max_row + 1):
            m = re.match(r"=\+PROYECTOS!B(\d+)$", str(ws.cell(r, 2).value))
            if m and int(m.group(1)) != r:
                bad += 1
        print(f"  {sh:22} filas desalineadas: {bad}")
    print("  Vínculos externos:", len(sup_f._external_links))

    print("\n== 6. Cobros 2026: CASH FLOW vs INGRESOS TOTALES (ene-dic 2026) ==")
    pr = sup_v["PROYECTOS"]
    it = sup_v["INGRESOS TOTALES"]
    filas = {}
    for r in range(1, pr.max_row + 1):
        d = pr.cell(r, 4).value
        if d:
            filas.setdefault(str(d).strip(), []).append(r)
    ok = dif = 0
    for r in range(4, 104):
        cod = cfv.cell(r, 3).value
        if not cod:
            continue
        cf_ing = sum(num(cfv.cell(r, k).value) for k in range(41, 45))
        sd_ing = sum(num(it.cell(R, k).value) for R in filas.get(str(cod).strip(), [])
                     for k in range(ci("BB"), ci("BM") + 1))
        if abs(cf_ing - sd_ing) < 1:
            ok += 1
        else:
            dif += 1
            print(f"  {cod}: CASH FLOW {cf_ing:>12,.2f}  SuperDoc {sd_ing:>12,.2f}")
    print(f"  Coinciden {ok} códigos, difieren {dif}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
